"""Tests for tools/xlsx (run: python3 -I -m unittest discover -s tests/py). Rules text here is made up."""
import importlib.util
import json
import os
import subprocess
import sys
import tempfile
import unittest

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
TOOLS = os.path.join(ROOT, 'tools', 'xlsx')


def load(name):
    spec = importlib.util.spec_from_file_location(name, os.path.join(TOOLS, name + '.py'))
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


try:
    import openpyxl  # noqa: F401  (both tools need it)
    mechanics, build = load('mechanics'), load('build')
except ImportError:
    mechanics = build = None


@unittest.skipIf(mechanics is None, 'openpyxl is not installed')
class Meta(unittest.TestCase):
    def test_parse_and_format_round_trip(self):
        d = mechanics.parse_meta('roll=Snark gain="2 luck" start')
        self.assertEqual(d, {'roll': 'Snark', 'gain': '2 luck', 'start': True})
        self.assertEqual(mechanics.parse_meta(mechanics.format_meta(d)), d)

    def test_merge_keeps_old_keys_and_replaces_named_ones(self):
        self.assertEqual(mechanics.merge_meta('roll=Snark start', 'roll=Grit spend=1'), 'roll=Grit start spend=1')


@unittest.skipIf(build is None, 'openpyxl is not installed')
class Overlay(unittest.TestCase):
    def test_theme_rows_fill_override_columns_or_are_added(self):
        book = [{'kind': 'move', 'group': 'basic', 'name': 'POKE', 'text': 'Poke.', 'meta': ''}]
        theme = [
            {'kind': 'move', 'group': 'basic', 'name': 'POKE', 'override_name': 'JAB'},
            {'kind': 'move', 'group': 'basic', 'name': 'MISSING', 'override_text': 'Only an override.'},
            {'kind': 'move', 'group': 'basic', 'name': 'NEW', 'text': 'A new move.'},
        ]
        rows, loose = build.overlay(book, theme)
        self.assertEqual(rows[0]['override_name'], 'JAB')
        self.assertEqual([r['name'] for r in rows], ['POKE', 'NEW'])
        self.assertEqual([t['name'] for t in loose], ['MISSING'])


@unittest.skipIf(mechanics is None, 'openpyxl is not installed')
class Workbook(unittest.TestCase):
    def run_tool(self, *args):
        subprocess.run([sys.executable, '-I', *args], check=True, cwd=ROOT, capture_output=True)

    def test_kind_dropdowns_go_on_character_and_ship_tabs_only(self):
        with tempfile.TemporaryDirectory() as d:
            p = lambda n: os.path.join(d, n)

            def write(name, data):
                with open(p(name), 'w', encoding='utf-8') as f:
                    json.dump(data, f)
            write('r.json', [{'kind': 'move', 'group': 'basic', 'name': 'POKE', 'text': 'Poke.', 'meta': '', 'vis': 'crew'}])
            write('n.json', [['info', 'Name', 'Nexa', '']])
            write('m.json', {'rows': [], 'meta': []})
            self.run_tool('tools/xlsx/build.py', '--rules', p('r.json'), '--tab', 'Nexa=' + p('n.json'), '--out', p('a.xlsx'))
            wb = openpyxl.load_workbook(p('a.xlsx'))
            for name in ('Ship', 'Feedback', 'Images'):
                if name not in wb.sheetnames:
                    wb.create_sheet(name).append(['Kind', 'Name', 'Value', 'Notes'])
            wb.save(p('a.xlsx'))
            self.run_tool('tools/xlsx/mechanics.py', '--mechanics', p('m.json'), p('a.xlsx'), '--out', p('b.xlsx'))
            wb = openpyxl.load_workbook(p('b.xlsx'))
            kinds = lambda name: [v.formula1 for v in wb[name].data_validations.dataValidation]
            self.assertIn('jot', kinds('Nexa')[0])
            self.assertTrue(kinds('Ship'))
            self.assertEqual(kinds('Feedback'), [])
            self.assertEqual(kinds('Images'), [])
            # The schema build.py writes outlasts mechanics.py (the import reads it).
            self.assertEqual(wb.custom_doc_props['Cybersea schema'].value, build.SCHEMA)


if __name__ == '__main__':
    unittest.main()

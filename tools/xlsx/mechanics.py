#!/usr/bin/env python3
"""Add a game's mechanics to a game workbook (.xlsx), and the checks that guard it.

The site has no stats, tracks or numbers of its own: they are rows in the System tab, and what a
Rules row rolls or does is in its Meta (see assets/system.js and SETUP.md). This lays them into an
existing workbook, to bring into a game's sheet with Cybersea > Import an xlsx…:

  python3 -I tools/xlsx/mechanics.py --mechanics mechanics.json in.xlsx --out out.xlsx

mechanics.json: {rows, meta}
  rows: the System tab's rows, as {kind, group, name, text, meta}. The tab is made anew.
  meta: [{kind, group, name, meta}]: keys added to that Rules row's Meta, key by key (roll=…, action=…).

It also adds the dropdowns Code.gs's setUpChecks makes on the Kind columns (and Rules' Vis), so
the workbook has them in Excel too. Import an xlsx… adds the rest as it imports: the Value
checks and the warning-only protection on Rules, System and README.
"""
import argparse
import json
import re

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

RULES_KINDS = ['move', 'procedure', 'house', 'advance', 'playbook', 'vice', 'oddity', 'kit', 'lastwords', 'question',
               'pmove', 'smove', 'pnote', 'renown', 'rule', 'special', 'tag', 'trouble', 'table', 'fate', 'world',
               'faction']
SYSTEM_KINDS = ['stat', 'track', 'section', 'choice']
SYSTEM_HEADER = ['Kind', 'Group', 'Name', 'Text', 'Meta']
SHEET_KINDS = ['info', 'stat', 'track', 'bond', 'sheet', 'notes', 'jot', 'kit', 'item', 'cargo', 'crew', 'move', 'special', 'advance', 'log', 'state', 'signin']
# As in Code.gs: tabs with these names are that kind; any other is a character (tests/kinds.test.mjs).
TAB_KINDS = ['rules', 'system', 'gm', 'lore', 'ship', 'map', 'feedback', 'images']
META_RE = re.compile(r'([^\s=]+)=(?:"([^"]*)"|(\S*))|(\S+)')


def parse_meta(s):
    out = {}
    for m in META_RE.finditer(str(s or '')):
        if m.group(4):
            out[m.group(4).lower()] = True
        else:
            out[m.group(1).lower()] = m.group(2) if m.group(2) is not None else m.group(3)
    return out


def format_meta(d):
    def one(k, v):
        if v is True:
            return k
        v = str(v)
        return f'{k}="{v}"' if re.search(r'[\s"]', v) or v == '' else f'{k}={v}'
    return ' '.join(one(k, v) for k, v in d.items())


def merge_meta(old, add):
    d = parse_meta(old)
    d.update(parse_meta(add))
    return format_meta(d)


def key(kind, group, name):
    return tuple(str(x or '').strip().lower() for x in (kind, group, name))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--mechanics', required=True)
    ap.add_argument('workbook')
    ap.add_argument('--out', required=True)
    a = ap.parse_args()
    mech = json.load(open(a.mechanics, encoding='utf-8'))
    wb = load_workbook(a.workbook)
    ws = wb['Rules']
    head = [str(c.value or '').strip().lower() for c in ws[1]]
    col = {h: i + 1 for i, h in enumerate(head) if h}
    at = {}
    last = 1
    for r in range(2, ws.max_row + 1):
        k = key(*(ws.cell(r, col[h]).value for h in ('kind', 'group', 'name')))
        if any(k):
            at.setdefault(k, r)
            last = r

    # The System tab, made anew after Rules; any system rows left in Rules are cleared.
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(r, col['kind']).value or '').strip().lower() in SYSTEM_KINDS:
            for c in range(1, ws.max_column + 1):
                ws.cell(r, c).value = None
    if 'System' in wb.sheetnames:
        del wb['System']
    sy = wb.create_sheet('System', wb.sheetnames.index('Rules') + 1)
    sy.append(SYSTEM_HEADER)
    rows = mech.get('rows', [])
    for row in rows:
        sy.append([str(row.get(h.lower(), '')) for h in SYSTEM_HEADER])
    for line in sy.iter_rows(min_row=1, max_row=1000, max_col=len(SYSTEM_HEADER)):
        for c in line:
            c.number_format = '@'
            c.alignment = Alignment(wrap_text=True, vertical='top')
    for c in sy[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill('solid', fgColor='E8EEF3')
    sy.freeze_panes = 'A2'
    for letter_, w in zip('ABCDE', [10, 16, 22, 40, 50]):
        sy.column_dimensions[letter_].width = w
    sk = DataValidation(type='list', formula1='"' + ','.join(SYSTEM_KINDS) + '"', allow_blank=True,
                        showErrorMessage=True, errorTitle='Kind', error='A kind the site knows (see README).')
    sy.add_data_validation(sk)
    sk.add('A2:A1000')

    tagged = 0
    missing = []
    for m in mech.get('meta', []):
        r = at.get(key(m.get('kind'), m.get('group'), m.get('name')))
        if not r:
            missing.append(f"{m.get('kind')} | {m.get('group')} | {m.get('name')}")
            continue
        c = ws.cell(r, col['meta'])
        c.value = merge_meta(c.value, m.get('meta', ''))
        c.number_format = '@'
        tagged += 1

    def letter(n):
        return ws.cell(1, n).column_letter

    # Checks (Code.gs setUpChecks makes the same ones in Google Sheets).
    ws.data_validations.dataValidation = []
    kinds = DataValidation(type='list', formula1='"' + ','.join(RULES_KINDS) + '"', allow_blank=True,
                           showErrorMessage=True, errorTitle='Kind', error='A kind the site knows (see README).')
    ws.add_data_validation(kinds)
    kinds.add(f"{letter(col['kind'])}2:{letter(col['kind'])}1000")
    vis = DataValidation(type='list', formula1='"crew,gm"', allow_blank=True, showErrorMessage=True,
                         errorTitle='Vis', error='crew (or blank): everyone sees it. gm: only the GM.')
    ws.add_data_validation(vis)
    vis.add(f"{letter(col['vis'])}2:{letter(col['vis'])}1000")

    for sh in wb.worksheets:
        name = sh.title.strip().lower()
        if name == 'readme' or name.startswith('_') or (name in TAB_KINDS and name != 'ship'):
            continue
        sh.data_validations.dataValidation = []
        dv = DataValidation(type='list', formula1='"' + ','.join(SHEET_KINDS) + '"', allow_blank=True, showErrorMessage=False)
        sh.add_data_validation(dv)
        dv.add('A2:A1000')

    wb.save(a.out)
    print(f'wrote {a.out}: System tab with {len(rows)} rows; {tagged} Rules rows tagged')
    for m in missing:
        print(f'  no row for: {m}')


if __name__ == '__main__':
    main()
